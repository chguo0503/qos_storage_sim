"""Editable presentation of existing mixed-input experiments; never run simulations.

Visual conventions follow coflow_global_5ms_experiments/presentation.
All metrics come from the final audited CSVs; input tables and sequence cells
are independently read from the frozen request manifests.
"""
from __future__ import annotations

import ast
import csv
import gzip
import hashlib
import json
import math
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parents[2]
ROOT = BASE.parents[1]
OUT = BASE / "presentation"
FONT = "Noto Sans CJK SC"
INK, MUTED, HEADER, GRID = "222222", "666666", "F0F1F3", "D6D6D6"
TEAL, AMBER = "007D78", "A85A22"
PALETTE = ["6196B8", "7FB7AE", "A5B98A", "C8A16A", "696E9B"]
POLICIES = [
    ("baseline", "固定单路", "Baseline / fixed", "每盘读取固定进入 Path0", ""),
    ("once", "压力感知选路", "Once / fixed", "根据 QoS Path 的 I/O 压力", "为每一层选择路径"),
    ("new_once", "预留感知选路", "New Once / fixed", "在 Once 基础上，计入全局", "已安排但未完成的 I/O"),
    ("strategy1_pipeline", "全局客户端调度", "策略 1 / pipeline", "选择 NPU，并控制", "I/O 下发时机"),
    ("strategy2_pipeline", "盘内协同调度", "策略 2 / pipeline", "增加盘内 I/O 调度", "保留原生 QoS 约束"),
    ("strategy3_pipeline", "跨盘协同调度", "策略 3 / pipeline", "结合各盘同层读取进度", "协调 I/O 优先级"),
    ("strategy3_fixed", "跨盘协同，固定分卡", "策略 3 / fixed", "保留输入的 NPU 绑定", "使用策略 3 的 I/O 调度"),
]
CASES = {
    "V38": ("mixed_varied/inputs/raw_size_varied_rho097_seed7.json.gz", "原始参数混合", "raw_size_varied_rho097_seed7"),
    "E72": ("mixed_varied/inputs/aligned_short072_seed7.json.gz", "外推短流混合", "aligned_short072_seed7"),
    "A6": ("mixed_six_ssu/inputs/aligned_short080_ssu6_seed7.json.gz", "六盘外推短流混合", "aligned_short080_ssu6_seed7"),
    "B6": ("mixed_six_ssu/inputs/raw_size_varied_ssu6_seed7.json.gz", "六盘原始参数混合", "raw_size_varied_ssu6_seed7"),
}
SOURCES = {}
DISPLAYED = []


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source(path):
    path = Path(path)
    SOURCES[str(path.relative_to(ROOT))] = sha(path)
    return path


def read(path):
    path = source(path)
    with (gzip.open(path, "rt") if path.suffix == ".gz" else path.open()) as f:
        return json.load(f)


def csv_read(path):
    with source(path).open() as f:
        return list(csv.DictReader(f))


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")


def rgb(value):
    return RGBColor.from_string(value)


def shape(slide, x, y, w, h, fill):
    obj = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    obj.line.fill.background()
    obj._element.spPr.append(OxmlElement("a:effectLst"))
    style = obj._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        obj._element.remove(style)
    return obj


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT):
    value = str(value)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name, run.font.size = font, Pt(size)
    run.font.bold, run.font.color.rgb = bold, rgb(color)
    props = run._r.get_or_add_rPr()
    props.set("lang", "zh-CN")
    ea = OxmlElement("a:ea")
    ea.set("typeface", FONT)
    props.append(ea)
    return box


def new_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")
    text(slide, .65, .47, 14.7, .60, title, 28, bold=True)
    text(slide, .67, 1.16, 14.6, .30, subtitle, 13.5, MUTED)
    return slide


def footer(slide, lines, page):
    for i, line in enumerate(lines):
        text(slide, .67, 8.02 + i * .28, 14.3, .24, line, 10.5, MUTED)
    text(slide, 15.05, 8.60, .30, .18, str(page), 9, MUTED, align=PP_ALIGN.RIGHT)


def table(slide, x, y, widths, headers, rows, row_h=.62, header_h=.62, size=17,
          header_size=14, highlights=None):
    total = sum(widths)
    shape(slide, x, y, total, header_h, HEADER)
    xx = x
    for j, (width, name) in enumerate(zip(widths, headers)):
        text(slide, xx + .12, y + .09, width - .24, header_h - .18,
             name, header_size, bold=True, align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
        xx += width
    for i, row in enumerate(rows):
        yy, xx = y + header_h + i * row_h, x
        for j, (width, value) in enumerate(zip(widths, row)):
            color, bold = (highlights or {}).get((i, j), (INK, False))
            text(slide, xx + .12, yy + .08, width - .24, row_h - .16,
                 value, size, color, bold,
                 PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER,
                 font=FONT if j == 0 else "DejaVu Sans")
            xx += width
        shape(slide, x, yy + row_h - .005, total, .005, GRID)
    xx = x
    for width in widths[:-1]:
        xx += width
        shape(slide, xx, y, .004, header_h + row_h * len(rows), "DDDDDD")


def profile_name(p, category=False):
    name = f"{p['seq_len_k']}K / {p['nql']}"
    return name + (f" · {p['category']}" if category else "")


def load_data():
    verification = read(BASE / "final_verification.json")
    assert verification["all_checks_pass"]
    frozen = {r["path"]: r["sha256"] for r in verification["audited_input_files"]}
    raw = ast.literal_eval(source(ROOT / "data").read_text())
    cases = {}
    for key, (relative, name, label) in CASES.items():
        path = BASE / relative
        assert sha(path) == frozen[str(path.relative_to(ROOT))]
        m = read(path)
        meta = m["metadata"]
        assert meta["num_npu"] == 32 and meta["n_layers"] == 8 and meta["seed"] == 7
        assert meta["compute_scale_actual"] == 1
        lanes = {n: [r for r in m["requests"] if r["npu_id"] == n] for n in range(32)}
        keys = [(int(p["seq_len_k"]), int(p["nql"])) for p in meta["profiles"]]
        profiles = []
        total_c = sum(r["load"]["per_layer_us"] * 8 / 1000 for r in lanes[0])
        for i, k in enumerate(keys):
            sample = next(r for r in lanes[0] if (r["load"]["seq_len_k"], r["load"]["nql"]) == k)
            load = sample["load"]
            per_card = sum((r["load"]["seq_len_k"], r["load"]["nql"]) == k for r in lanes[0])
            counts = [sum((r["load"]["seq_len_k"], r["load"]["nql"]) == k for r in lane) for lane in lanes.values()]
            assert counts == [per_card] * 32
            placement = m["placements"][sample["placement_index"]]
            assert len(placement) == 1
            v = sum(block[1] for block in placement[0])
            c = load["per_layer_us"] / 1000
            assert math.isclose(v, load["per_layer_kv_gb"], abs_tol=1e-10)
            assert math.isclose(v * 2**30, (k[0] * 1024 - k[1]) * 1408, abs_tol=1e-6)
            assert load["padding_gib_per_layer"] == 0
            direct = load["profile_construction"]["method"] == "direct_data_row"
            if direct:
                assert math.isclose(raw[k][1] / 1000, c, rel_tol=1e-12)
                assert math.isclose(raw[k][3], v, abs_tol=1e-10)
            profiles.append({"seq_len_k": k[0], "nql": k[1], "category": load["category"],
                             "role": load["role"], "C_ms_per_layer": c, "V_MiB_per_layer": v * 1024,
                             "requests_per_npu": per_card, "requests_total": per_card * 32,
                             "request_fraction": per_card / len(lanes[0]),
                             "compute_fraction": per_card * c * 8 / total_c,
                             "direct_data_row": direct, "color": PALETTE[i]})
        sequences = [[keys.index((r["load"]["seq_len_k"], r["load"]["nql"])) for r in lane]
                     for lane in lanes.values()]
        assert len({tuple(lane) for lane in sequences}) == 32
        assert all(r["arrival_time_ms"] == 0 for r in m["requests"])
        assert meta["profile_counts_per_npu"] == [p["requests_per_npu"] for p in profiles]
        demand = meta["input_demand"]
        cases[key] = {"key": key, "label": label, "name": name, "manifest": relative,
                      "input_fingerprint": m["input_fingerprint"], "num_ssu": meta["num_ssu"],
                      "requests_per_npu": len(lanes[0]), "requests_total": len(m["requests"]),
                      "total_C_ms_per_npu": total_c,
                      "hottest_ssu_rho": demand["hottest_ssu_load_ratio"],
                      "largest_receive_link_rho": demand["largest_npu_receive_load_ratio"],
                      "short_C_fraction": sum(p["compute_fraction"] for p in profiles if p["role"] == "short"),
                      "profiles": profiles, "sequences": sequences,
                      "request_ids": [[r["request_id"] for r in lane] for lane in lanes.values()]}
    runs = csv_read(BASE / "notes/mixed_policy_audit/runs.csv")
    assert len(runs) == 14 and all(r["status"] == "complete" for r in runs)
    run_index = {(r["label"], r["strategy"]): r for r in runs}
    for key in ("E72", "V38"):
        for policy, *_ in POLICIES:
            row = run_index[cases[key]["label"], policy]
            assert row["input_fingerprint"] == cases[key]["input_fingerprint"]
            assert row["all_active_1"] == row["all_active_2"] == "True"
    six = csv_read(BASE / "mixed_six_ssu/analysis/windows.csv")
    profiles = csv_read(BASE / "notes/mixed_policy_audit/profiles.csv")
    return cases, run_index, six, profiles


def first_slide(prs, cases, runs):
    slide = new_slide(prs, "混合请求下的 QoS 策略对比", "32 NPU / 8 SSU，单盘 40 GiB/s，状态采集周期 5 ms；每张卡输入多种请求")
    x, y, total, sh, ch = .65, 1.77, 14.70, 2.85, 4.05
    header_h, row_h = .78, .65
    nx, group_w = x + sh + ch, (total - sh - ch) / 2
    shape(slide, x, y, total, header_h, HEADER)
    text(slide, x + .16, y + .20, sh - .3, .35, "策略", 16, bold=True)
    text(slide, x + sh + .12, y + .20, ch - .25, .35, "主要能力", 16, bold=True)
    values = {}
    for j, key in enumerate(("E72", "V38")):
        xx = nx + j * group_w
        text(slide, xx, y + .075, group_w, .30, f"{cases[key]['name']} · {key}", 15, bold=True, align=PP_ALIGN.CENTER)
        for k, name in enumerate(("利用率 %", "SLO 达标 %")):
            text(slide, xx + k * group_w / 2, y + .47, group_w / 2, .22, name, 11.5, MUTED, align=PP_ALIGN.CENTER)
        for policy, *_ in POLICIES:
            row = runs[cases[key]["label"], policy]
            values[policy, key] = [100 * float(row["U1"]), 100 * float(row["all_SLO"])]
    maxima = {(key, k): max(round(values[p, key][k], 2) for p, *_ in POLICIES)
              for key in ("E72", "V38") for k in (0, 1)}
    for i, (policy, label, old, c1, c2) in enumerate(POLICIES):
        yy = y + header_h + i * row_h
        text(slide, x + .16, yy + .08, sh - .24, .29, label, 16)
        text(slide, x + .17, yy + .40, sh - .25, .18, old, 10, MUTED)
        if c2:
            text(slide, x + sh + .12, yy + .09, ch - .24, .24, c1, 13.5)
            text(slide, x + sh + .12, yy + .37, ch - .24, .24, c2, 13.5)
        else:
            text(slide, x + sh + .12, yy + .20, ch - .24, .25, c1, 13.5)
        for j, key in enumerate(("E72", "V38")):
            for k, val in enumerate(values[policy, key]):
                text(slide, nx + (j + k / 2) * group_w, yy + .15, group_w / 2, .34,
                     f"{val:.2f}", 18, bold=round(val, 2) == maxima[key, k], align=PP_ALIGN.CENTER, font="DejaVu Sans")
                DISPLAYED.append({"slide": 1, "case": key, "policy": policy, "metric": ("U1_percent", "all_admission_SLO_percent")[k], "value": val})
        shape(slide, x, yy + row_h - .005, total, .005, GRID)
    for xx in (x + sh, nx, nx + group_w):
        shape(slide, xx, y, .004, header_h + len(POLICIES) * row_h, "DDDDDD")
    text(slide, .67, 7.37, 14.55, .34, "整机利用率高，仍可能有短请求等待；选路、选卡与整批完成时间需要分别评价。", 15)
    footer(slide, [
        "注：利用率取 [1,2] 秒内32卡平均；SLO覆盖完整输入，接纳至8层完成 ≤ 1.5 × 纯计算时间，不含接纳前排队。",
        "seed7；E72共18,624条，V38共2,944条。粗体为列内显示最高值。E72含1K外推；V38画像参数来自data。",
    ], 1)
    return slide


def weight_bar(slide, case, y):
    text(slide, .68, y, 14.6, .30, "完整输入的纯计算时间构成", 14, bold=True)
    xx = .65
    for p in case["profiles"]:
        width = 14.7 * p["compute_fraction"]
        shape(slide, xx, y + .45, width, .36, p["color"])
        if width > .8:
            text(slide, xx, y + .47, width, .30, f"{100*p['compute_fraction']:.1f}%", 11, "FFFFFF" if p["color"] == PALETTE[-1] else INK,
                 align=PP_ALIGN.CENTER, font="DejaVu Sans")
        xx += width
    width = 14.7 / len(case["profiles"])
    for i, p in enumerate(case["profiles"]):
        xx = .68 + i * width
        shape(slide, xx, y + .98, .13, .13, p["color"])
        text(slide, xx + .21, y + .90, width - .26, .27, profile_name(p), 11, MUTED)


def input_slide(prs, case, runs, page):
    raw = case["key"] == "V38"
    title = "输入一：整机利用率约96%的请求混合" if raw else "输入二：整机利用率约80%的请求混合"
    slide = new_slide(prs, title,
                      f"{case['key']}，32 NPU / 8 SSU；每卡{case['requests_per_npu']}条，全局{case['requests_total']:,}条；各卡配额相同、完整顺序独立")
    b, n = [runs[case["label"], p] for p in ("baseline", "new_once")]
    text(slide, .67, 1.70, 9.9, .34,
         f"Baseline U1/U2：{100*float(b['U1']):.2f}% / {100*float(b['U2']):.2f}%     New Once：{100*float(n['U1']):.2f}% / {100*float(n['U2']):.2f}%", 16)
    text(slide, 11.65, 1.70, 3.7, .34, f"最热盘平均ρ：{100*case['hottest_ssu_rho']:.2f}%", 14, MUTED, align=PP_ALIGN.RIGHT)
    rows = [[profile_name(p, True), f"{p['C_ms_per_layer']:.3f}", f"{p['V_MiB_per_layer']:.3f}",
             str(p["requests_per_npu"]), f"{100*p['request_fraction']:.2f}%", f"{100*p['compute_fraction']:.2f}%"] for p in case["profiles"]]
    table(slide, .65, 2.25, [2.9, 1.85, 2.35, 1.75, 2.8, 3.05],
          ["序列长度 / NQL · 类别", "每层 C (ms)", "每层读取 (MiB)", "每卡条数", "请求数量占比", "纯计算时间占比"], rows,
          row_h=.59, header_h=.62, size=17, header_size=12.8)
    if raw:
        line = "每卡28+28+28条较短请求、2+6条长请求；192K两类合计占62.45%的纯计算时间。"
    else:
        line = "每卡188+188+188条低读取量短请求、18条大读取请求；三短画像占71.82%的纯计算时间。"
    text(slide, .67, 6.02, 14.6, .34, line, 15)
    weight_bar(slide, case, 6.51)
    footer(slide, [
        "K=1024 token；NQL为本次新增query token数。C为每层纯计算时间；读取量为SSD→HBM的跨盘合计。",
        "每请求8层、batch=1；所有请求t=0到达。" + ("五种画像均取原data；无C缩放、无尾块补齐。" if raw else "三个1K画像使用外推，NQL384还含插值；192K画像取原data，无C缩放。"),
    ], page)
    return slide


def six_slide(prs, cases, six):
    slide = new_slide(prs, "32 NPU / 6 SSU：较低利用率对应什么输入", "两组不同请求配额；每组都在自己的冻结输入上比较Baseline与New Once，seed7")
    for j, key in enumerate(("A6", "B6")):
        c = cases[key]
        x, width = .65 + j * 7.65, 7.05
        text(slide, x, 1.75, width, .36, "A：三种外推短流 + 一种长流" if j == 0 else "B：三种原始参数短流 + 两种长流", 18, bold=True)
        for i, policy in enumerate(("baseline", "new_once")):
            matching = [r for r in six if r["label"] == c["label"] and r["strategy"] == policy]
            assert len(matching) == 2
            matching.sort(key=lambda r: float(r["window_start_ms"]))
            vals = [100 * float(r["utilization"]) for r in matching]
            text(slide, x, 2.25 + i * .38, width, .30,
                 f"{'Baseline' if i == 0 else 'New Once'} U1/U2：{vals[0]:.2f}% / {vals[1]:.2f}%", 17, INK if i == 0 else TEAL)
            for k, value in enumerate(vals):
                DISPLAYED.append({"slide": 4, "case": key, "policy": policy, "metric": f"U{k+1}_percent", "value": value})
        rows = [[profile_name(p), str(p["requests_per_npu"]), f"{100*p['compute_fraction']:.2f}%"] for p in c["profiles"]]
        table(slide, x, 3.19, [3.10, 1.70, 2.25], ["序列长度 / NQL", "每卡条数", "C时间占比"], rows,
              row_h=.48, header_h=.53, size=16, header_size=13)
        text(slide, x, 6.38, width, .30, f"每卡{c['requests_per_npu']}条；全局{c['requests_total']:,}条", 16)
        text(slide, x, 6.84, width, .30,
             f"最热盘平均ρ {100*c['hottest_ssu_rho']:.2f}%；短组C占比 {100*c['short_C_fraction']:.2f}%", 14, MUTED)
    text(slide, .67, 7.42, 14.6, .34, "A的短请求占据更多计算时间，Baseline整体损失更明显；B仍需逐类查看短请求等待。", 15)
    footer(slide, [
        "U1=[1,2]秒，U2=[2,3]秒；全部32卡全窗active。每盘40GiB/s、每卡接收50GiB/s；平均ρ不保证层期限可满足。",
        "A对应8盘E80的同人口对照，区别于前页E72。B在8盘V38基础上重新配额；两者都无C缩放。",
    ], 4)
    return slide


def sequence_slide(prs, case):
    slide = new_slide(prs, "每张NPU究竟收到什么顺序的请求？", "V38：每张卡92条，下图绘制全部2,944条输入请求；每格等宽表示一条请求的位置")
    for i, p in enumerate(case["profiles"]):
        x = .68 + i * 2.94
        shape(slide, x, 1.77, .16, .16, p["color"])
        text(slide, x + .25, 1.70, 2.62, .30, profile_name(p), 12)
    x, y, width, row_h = 1.36, 2.48, 13.99, .127
    cw = width / case["requests_per_npu"]
    text(slide, .67, 2.11, .62, .24, "NPU", 10, MUTED)
    for k in [1] + list(range(10, 91, 10)) + [92]:
        text(slide, x + (k - .5) * cw - .13, 2.13, .26, .22, str(k), 9, MUTED, align=PP_ALIGN.CENTER, font="DejaVu Sans")
    for n, lane in enumerate(case["sequences"]):
        yy = y + n * row_h
        text(slide, .74, yy, .48, row_h, f"{n:02d}", 8.3, MUTED, align=PP_ALIGN.RIGHT, font="DejaVu Sans")
        for position, profile in enumerate(lane):
            shape(slide, x + position * cw, yy, cw - .006, row_h - .018, case["profiles"][profile]["color"])
    text(slide, .67, 6.92, 14.6, .34, "相同配额，各卡独立乱序；相同画像会多次出现，但没有重复固定的小循环。", 17)
    text(slide, .67, 7.46, 14.6, .30, "完整列表一次shuffle：Random(seed + NPU编号 × 100003)，本页seed=7；32条完整顺序各不相同。", 13.5, MUTED)
    footer(slide, [
        "这是输入位置图，各格的实际计算时长不同。所有请求t=0入队，固定分卡策略按各卡序列执行。",
        "S1–S3 pipeline还会重新分配NPU。完整原始输入和逐条顺序CSV随报告保存；频率与到达并非生产trace。",
    ], 5)
    return slide


def cohort_slide(prs, case, profiles):
    slide = new_slide(prs, "96%的整机平均，掩盖了哪些请求？", "相同V38种子7：整机U取[1,2]秒；下表逐类指标覆盖同一批完整请求，保留长流和变差的画像")
    text(slide, .67, 1.74, 14.6, .34, "逐类计算占比 = 该类纯计算总时间 / 该类接纳至完成总历时", 16)
    ix = {(r["strategy"], int(r["seq_len_k"]), int(r["nql"])): r for r in profiles if r["label"] == case["label"]}
    rows, highlights = [], {}
    for i, p in enumerate(case["profiles"]):
        b, n = [ix[policy, p["seq_len_k"], p["nql"]] for policy in ("baseline", "new_once")]
        bu, nu = [100 * float(r["aggregate_compute_fraction"]) for r in (b, n)]
        rows.append([profile_name(p, True), f"{100*p['compute_fraction']:.2f}%", f"{bu:.2f}%", f"{nu:.2f}%",
                     f"{b['admission_slo_passed']}/{b['request_count']}", f"{n['admission_slo_passed']}/{n['request_count']}"])
        highlights[i, 3] = (TEAL if nu >= bu else AMBER, False)
        for policy, r, value in (("baseline", b, bu), ("new_once", n, nu)):
            DISPLAYED.append({"slide": 6, "case": "V38", "policy": policy, "profile": [p["seq_len_k"], p["nql"]],
                              "metric": "full_cohort_compute_fraction_percent", "value": value,
                              "admission_slo_passed": int(r["admission_slo_passed"]), "request_count": int(r["request_count"])})
    table(slide, .65, 2.34, [2.90, 1.90, 2.50, 2.50, 2.45, 2.45],
          ["序列长度 / NQL · 类别", "C时间权重", "Baseline计算占比", "New Once计算占比", "Baseline SLO", "New Once SLO"], rows,
          row_h=.61, header_h=.65, size=17, header_size=12.5, highlights=highlights)
    text(slide, .67, 6.46, 14.6, .36, "192K两类占62.45%的纯计算时间；最短32K/128仅占4.32%，其等待容易被整体平均掩盖。", 15)
    text(slide, .67, 7.13, 14.6, .37, "New Once更好地保护前两类，但64K/512和长类的计算占比下降；不能概括为所有请求都更快。", 15)
    footer(slide, [
        "SLO：接纳至完成 ≤ 1.5 × 自身8层纯计算时间；不含接纳前排队。64K/NQL512属于SL，短组并非全属SS。",
        "另一组三种严格SS的两种子复核也发现等待，详见完整研究报告；本页不把不同人口的结果混合平均。",
    ], 6)
    return slide


def export_input_docs(cases, runs):
    fields = ["case", "num_ssu", "seq_len_k", "nql", "category", "role", "C_ms_per_layer", "V_MiB_per_layer",
              "requests_per_npu", "requests_total", "request_fraction", "compute_fraction", "direct_data_row"]
    with (OUT / "input_profiles.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fields)
        writer.writeheader()
        for key, case in cases.items():
            for profile in case["profiles"]:
                row = {"case": key, "num_ssu": case["num_ssu"], **profile}
                writer.writerow({k: row[k] for k in fields})
    with (OUT / "input_sequences.csv").open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["case", "original_npu_id", "position_1based", "request_id", "seq_len_k", "nql", "arrival_ms"])
        for key, case in cases.items():
            for n, lane in enumerate(case["sequences"]):
                for pos, p in enumerate(lane):
                    profile = case["profiles"][p]
                    writer.writerow([key, n, pos + 1, case["request_ids"][n][pos], profile["seq_len_k"], profile["nql"], 0])
    lines = ["# 混合输入报告：输入明细", "", "全部数据读取已完成、已审计的实验，没有新增模拟。",
             "", "K=1024 token；NQL是本次新增query token数。历史前缀为seq×1024−NQL个token；每token每层1408字节，SSD直接读取到HBM。",
             "", "共同条件：32张NPU；每盘40GiB/s，每卡接收链路50GiB/s；每请求8层，batch=1；176KiB I/O命令；只预取下一层及已到达的下一请求L0。",
             "", "全部请求t=0到达。每卡含相同完整配额，用Random(7+npu_id*100003)独立打乱整个列表；不是生产trace，也不是每张卡持续重复单一画像。",
             "", "V表示每层跨全部SSD的总读取量，单请求八层读取量=8V，纯计算时间=8C。平均盘ρ按完整输入每卡总读取量/总纯C统计，不保证逐层deadline。", ""]
    for key, c in cases.items():
        lines += [f"## {key}：{c['name']}，32 NPU / {c['num_ssu']} SSU", "",
                  f"每卡{c['requests_per_npu']}条，全局{c['requests_total']:,}条；每卡总纯C {c['total_C_ms_per_npu']:.6f}ms；最热盘ρ {100*c['hottest_ssu_rho']:.6f}%，最大接收链路ρ {100*c['largest_receive_link_rho']:.6f}%。", "",
                  "|序列K / NQL|类别|每层C ms|每层读取MiB|每卡条数|全局条数|请求数占比|纯C占比|参数来源|",
                  "|---|---|---:|---:|---:|---:|---:|---:|---|"]
        for p in c["profiles"]:
            lines.append(f"|{p['seq_len_k']}K/{p['nql']}|{p['category']}|{p['C_ms_per_layer']:.6f}|{p['V_MiB_per_layer']:.6f}|{p['requests_per_npu']}|{p['requests_total']}|{100*p['request_fraction']:.4f}%|{100*p['compute_fraction']:.4f}%|{'原data' if p['direct_data_row'] else '1K外推；384另含NQL插值'}|")
        lines += ["", f"[冻结输入](../{c['manifest']})；fingerprint：`{c['input_fingerprint']}`。", ""]
        for n in (0, 1):
            sequence = " → ".join(profile_name(c["profiles"][p]) for p in c["sequences"][n][:12])
            lines += [f"NPU {n} 的前12条：{sequence}。", ""]
    lines += ["## 不能混为同输入的比较", "", "E72八盘与A6六盘配额不同。A6真正的同人口八盘参照是E80（aligned_short080_seed7）。B6与八盘V38也重新配额，因此不能把两者差值单独归因于盘数。",
              "", "所有策略结果使用各自输入的同一冻结人口与物理位置。S1–S3 pipeline允许重分NPU；baseline、Once、New once及S3 fixed保留原绑定。",
              "", "[完整逐条输入顺序CSV](input_sequences.csv) · [画像参数CSV](input_profiles.csv) · [完整研究报告](../report.pdf)", ""]
    (OUT / "input_details.md").write_text("\n".join(lines))


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    cases, runs, six, profiles = load_data()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    prs.core_properties.title = "混合输入下的NPU利用率与QoS策略"
    prs.core_properties.subject = "已有实验的六页可编辑报告：策略对照、实际输入与请求服务质量"
    prs.core_properties.author = "QoS Storage Simulation"
    slides = [first_slide(prs, cases, runs), input_slide(prs, cases["V38"], runs, 2),
              input_slide(prs, cases["E72"], runs, 3), six_slide(prs, cases, six),
              sequence_slide(prs, cases["V38"]), cohort_slide(prs, cases["V38"], profiles)]
    provenance = {"sources": SOURCES, "method": "Read existing final audited results; no simulations; single seed7 unless explicitly stated.",
                  "input_scope": "Complete per-NPU populations independently shuffled once; all arrivals t=0; synthetic frequencies/order.",
                  "metric_scope": "Fleet U1=[1000,2000]ms, U2=[2000,3000]ms. All-request admission SLO and per-profile cohort ratios cover the complete population.",
                  "layer_units": "C and V tables are per layer; V sums all SSDs, 8 layers per request.",
                  "editability": "All tables, text, bars and 2944 input-sequence cells are native editable PowerPoint objects."}
    for i, slide in enumerate(slides, 1):
        slide.notes_slide.notes_text_frame.text = (
            f"第{i}页。\n" + json.dumps(provenance, ensure_ascii=False, indent=2) + "\n"
            + json.dumps([r for r in DISPLAYED if r["slide"] == i], ensure_ascii=False, indent=2)
        )
    target = OUT / "mixed_input_report.pptx"
    prs.save(target)
    check = Presentation(target)
    assert len(check.slides) == 6
    for slide in check.slides:
        assert not any(s.shape_type == 13 for s in slide.shapes), "Editable objects only"
        for s in slide.shapes:
            assert s.left >= 0 and s.top >= 0
            assert s.left + s.width <= check.slide_width + 2
            assert s.top + s.height <= check.slide_height + 2
    for name, digest in SOURCES.items():
        assert sha(ROOT / name) == digest, f"Source changed while building: {name}"
    export_input_docs(cases, runs)
    write_json(OUT / "source/presentation_data.json", {"cases": cases, "run_rows": list(runs.values()), "displayed_metrics": DISPLAYED, **provenance})
    write_json(OUT / "source/slide_audit.json", {
        "created_utc": datetime.now(timezone.utc).isoformat(), "slide_count": 6,
        "all_objects_editable": True, "geometry_within_slide": True,
        "source_hashes_verified_unchanged": True, "source_hashes": SOURCES,
        "builder_sha256": sha(Path(__file__)), "pptx_sha256": sha(target),
        "shape_counts": [len(s.shapes) for s in check.slides],
        "profile_counts_per_npu": {k: [p["requests_per_npu"] for p in c["profiles"]] for k, c in cases.items()},
        "displayed_metrics": DISPLAYED,
    })
    print(json.dumps({"pptx": str(target), "slides": len(check.slides), "requests_exported": sum(c["requests_total"] for c in cases.values())}, ensure_ascii=False))


if __name__ == "__main__":
    build()
