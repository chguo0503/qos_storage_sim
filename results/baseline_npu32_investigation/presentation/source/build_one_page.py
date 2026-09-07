"""Build the requested single editable slide from frozen experiment results."""
from datetime import datetime, timezone
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
import build_input_slides as b


def build():
    cases, runs, six, _ = b.load_data()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    prs.core_properties.title = "32 NPU：混合输入与不同SSU盘数下的策略利用率"
    slide = b.new_slide(prs, "32 NPU：混合输入下的策略利用率", "SSU盘数明确分列；单盘40 GiB/s，状态采集周期5 ms；表中数值为NPU利用率（%）")
    x, y, width, left, capability = .65, 1.72, 14.70, 2.50, 3.00
    cell = (width-left-capability)/4
    header, row_h = 1.00, .47
    b.shape(slide, x, y, width, header, b.HEADER)
    b.text(slide, x+.14, y+.32, left-.2, .3, "策略", 16, bold=True)
    b.text(slide, x+left+.12, y+.32, capability-.2, .3, "主要能力", 16, bold=True)
    start = x+left+capability
    for j, title in enumerate(("1K短请求 + 192K大读取请求", "32K / 48K / 64K + 192K请求")):
        b.text(slide, start+2*j*cell, y+.10, 2*cell, .3, title, 14, bold=True, align=PP_ALIGN.CENTER)
    order = ("A6", "E72", "B6", "V38")
    for j, key in enumerate(order):
        c=cases[key]
        b.text(slide, start+j*cell, y+.55, cell, .3,
               f"{c['num_ssu']}盘 SSU · {c['requests_per_npu']}条/卡", 12.5, bold=True, align=PP_ALIGN.CENTER)
    values = {}
    for policy, *_ in b.POLICIES:
        for key in order:
            if cases[key]['num_ssu']==8:
                value=100*float(runs[cases[key]['label'],policy]['U1'])
            else:
                rows=[r for r in six if r['label']==cases[key]['label'] and r['strategy']==policy and float(r['window_start_ms'])==1000]
                assert len(rows)<=1
                if rows:
                    assert rows[0]['all_active']=='True' and rows[0]['input_fingerprint']==cases[key]['input_fingerprint']
                value=100*float(rows[0]['utilization']) if rows else None
            values[policy,key]=value
    maxima={k:max(round(v,2) for (p,c),v in values.items() if c==k and v is not None) for k in order}
    names=["Baseline · 固定分卡", "Once · 固定分卡", "New Once · 固定分卡", "策略1 · 可重分NPU", "策略2 · 可重分NPU", "策略3 · 可重分NPU", "策略3 · 固定分卡"]
    caps=["每盘固定使用Path0", "按路径I/O压力选路", "选路计入全局未完成I/O", "全局选卡、控制I/O下发", "增加盘内I/O调度", "跨盘协调同层I/O优先级", "保留输入绑定，跨盘调度"]
    displayed=[]
    for i,(policy,*_) in enumerate(b.POLICIES):
        yy=y+header+i*row_h
        b.text(slide,x+.14,yy+.08,left-.22,.28,names[i],13)
        b.text(slide,x+left+.12,yy+.08,capability-.22,.28,caps[i],12.5)
        for j,key in enumerate(order):
            v=values[policy,key]
            s="—" if v is None else f"{v:.2f}"
            b.text(slide,start+j*cell,yy+.06,cell,.32,s,18,b.MUTED if v is None else b.INK,
                   bold=v is not None and round(v,2)==maxima[key],align=PP_ALIGN.CENTER)
            displayed.append({'policy':policy,'case':key,'num_ssu':cases[key]['num_ssu'],'U1_percent':v})
        b.shape(slide,x,yy+row_h-.004,width,.004,b.GRID)
    for xx in (x+left,start,start+cell,start+2*cell,start+3*cell):
        b.shape(slide,xx,y,.004,header+7*row_h,"DDDDDD")
    b.text(slide,.67,6.22,14.6,.3,"每张NPU的具体输入（序列长度 / NQL × 条数；K = 1024 token）",14,bold=True)
    recipes={
        'A6':("6盘：1K/128、1K/256、1K/384各195条；", "192K/256有12条。"),
        'E72':("8盘：1K/128、1K/256、1K/384各188条；", "192K/256有18条。"),
        'B6':("6盘：32K/128、48K/256、64K/512各18条；", "192K/1024有3条，192K/2048有9条。"),
        'V38':("8盘：32K/128、48K/256、64K/512各28条；", "192K/1024有2条，192K/2048有6条。"),
    }
    for j,pair in enumerate((('A6','E72'),('B6','V38'))):
        for i,key in enumerate(pair):
            yy=6.67+i*.64
            for line,t in enumerate(recipes[key]):
                b.text(slide,.67+j*7.48,yy+line*.25,7.15,.24,t,12.2)
    b.text(slide,.67,8.02,14.6,.24,"利用率：seed7，[1,2]秒内32卡平均；粗体为该列已测最高值；—为未测试。6盘与8盘配额不同，不能只归因于盘数。",10.5,b.MUTED)
    b.text(slide,.67,8.30,14.6,.24,"各卡配额相同、完整顺序独立打乱；全部t=0到达，每请求8层、batch=1。1K参数外推（NQL384含插值）；右组参数取原data。",10.5,b.MUTED)
    b.text(slide,.67,8.57,14.6,.23,"NQL为本次新增query token数。右组8盘中，192K两类占62.45%的纯计算时间，整机96.23%仍可能掩盖短请求等待。",10.2,b.MUTED)
    slide.notes_slide.notes_text_frame.text=json.dumps({'metrics':displayed,'sources':b.SOURCES,'input_profiles':{k:cases[k]['profiles'] for k in order},'metric':'seed7 fleet mean utilization [1000,2000] ms'},ensure_ascii=False,indent=2)
    target=b.OUT/'mixed_input_summary.pptx'
    prs.save(target)
    reread=Presentation(target)
    assert len(reread.slides)==1
    for s in reread.slides[0].shapes:
        assert s.shape_type!=13
        assert s.left>=0 and s.top>=0 and s.left+s.width<=prs.slide_width+2 and s.top+s.height<=prs.slide_height+2
    text=' '.join(s.text for s in reread.slides[0].shapes if s.has_text_frame)
    assert 'E72' not in text and 'V38' not in text and '6盘 SSU' in text and '8盘 SSU' in text
    assert all(b.sha(b.ROOT/name)==digest for name,digest in b.SOURCES.items())
    b.write_json(b.OUT/'source/one_page_audit.json',{'created_utc':datetime.now(timezone.utc).isoformat(),'slide_count':1,'editable':True,'source_hashes':b.SOURCES,'metrics':displayed,'pptx_sha256':b.sha(target),'builder_sha256':b.sha(Path(__file__))})
    print(str(target))


if __name__=='__main__':
    build()
