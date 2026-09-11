"""One editable, illustrative slide explaining non-fixed QoS Path mapping."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from build_strategy_slide import rgb, shape, text


OUT = Path(__file__).resolve().parents[1]
BLUE = "3569A4"
GRAY = "666666"
INK = "222222"
LINE = "B7BEC6"

# Only representative IOs are drawn. Both current requests are SS; Paths
# 0/1/2 are part of their permitted logical pool in the current static config.
# This is not an experimental trace or a diagram of physical hardware groups.
ROUTES = (
    ((0, 0, 0), (1, 0, 0), (0, 1, 1), (1, 1, 2)),
    ((0, 0, 2), (1, 0, 1), (0, 1, 0), (1, 1, 2)),
)
QUEUE_COUNTS = (
    ((1, 4, 4), (4, 1, 1)),
    ((4, 1, 1), (1, 4, 1)),
)


def arrow(slide, points, color, dashed=False):
    """Native editable line segments, with one arrow head at the destination."""
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        conn.line.color.rgb = rgb(color)
        conn.line.width = Pt(1.65)
        conn._element.spPr.append(OxmlElement("a:effectLst"))
        style = conn._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            conn._element.remove(style)
        if dashed:
            conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if i == len(points) - 2:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            end.set("w", "sm")
            end.set("len", "sm")
            conn.line._get_or_add_ln().append(end)


def panel(slide, x, stage):
    text(slide, x, 1.49, 6.76, .42,
         "第 L 层" if stage == 0 else "第 L+1 层", 20, INK, True)

    npu_x, npu_w = x + .08, 1.47
    npu_y = (2.91, 5.22)
    npu_h = .66
    asu_x, asu_w = x + 3.67, 2.96
    asu_y = (1.99, 4.79)
    path_x, path_w, path_h = asu_x + .18, asu_w - .36, .43
    targets = {}

    for npu in (0, 1):
        col = BLUE if npu == 0 else GRAY
        shape(slide, npu_x, npu_y[npu], npu_w, npu_h, "FFFFFF", line=col)
        text(slide, npu_x, npu_y[npu], npu_w, npu_h,
             f"NPU{npu}", 18, col, True, PP_ALIGN.CENTER)
        text(slide, npu_x, npu_y[npu] + npu_h + .10, npu_w, .27,
             "当前请求：SS", 11.5, GRAY, align=PP_ALIGN.CENTER)

    for asu in (0, 1):
        ay = asu_y[asu]
        shape(slide, asu_x, ay, asu_w, 2.50, "FFFFFF", line=LINE)
        text(slide, asu_x + .18, ay + .065, 1.1, .30, f"ASU{asu}", 16.5, INK, True)
        shape(slide, asu_x + .12, ay + .38, asu_w - .24, 1.76, "F8FAFC", line=LINE)
        text(slide, asu_x + .22, ay + .40, asu_w - .44, .24,
             "SS 类可选 Path（部分）", 11.5, INK, True)
        for path in range(3):
            py = ay + .69 + path * .49
            shape(slide, path_x, py, path_w, path_h, "FFFFFF", line=LINE)
            text(slide, path_x + .14, py, .74, path_h, f"Path{path}", 14.5, INK)
            for block in range(QUEUE_COUNTS[stage][asu][path]):
                bx, by, bw, bh = path_x + .98 + block * .38, py + .065, .32, .30
                io = shape(slide, bx, by, bw, bh, "E4E9EF", line="A2ADBA")
                io.name = f"queued_io_s{stage}_asu{asu}_path{path}_{block}"
                text(slide, bx, by, bw, bh, "IO", 9.5, INK, align=PP_ALIGN.CENTER)
            targets[asu, path] = (path_x, py + path_h / 2)
        shape(slide, asu_x + .12, ay + 2.20, asu_w - .24, .24, "EFEFEF")
        text(slide, asu_x + .12, ay + 2.20, asu_w - .24, .24,
             "其他类别：SL / LS / LL", 10.5, GRAY, align=PP_ALIGN.CENTER)

    # Separate elbows retain arrow ownership without a full connectivity mesh.
    for npu, asu, path in ROUTES[stage]:
        tx, ty = targets[asu, path]
        sx = npu_x + npu_w
        sy = npu_y[npu] + (.21 if asu == 0 else (.59 if npu == 1 else .46))
        if asu == 0 and npu == 0:
            elbow = x + 2.23
            if stage == 0:
                ty -= .07
        elif asu == 0 and npu == 1:
            elbow = x + 3.15
            if stage == 0:
                ty += .07
        elif asu == 1 and npu == 0:
            elbow = x + 1.96
        else:
            elbow = x + 2.72
        arrow(slide, ((sx, sy), (elbow, sy), (elbow, ty), (tx, ty)),
              BLUE if npu == 0 else GRAY, dashed=npu == 1)


def build():
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    prs.core_properties.title = "Once per layer：动态选择 QoS Path"
    prs.core_properties.subject = "按请求类别确定Path池，池内动态选路，NPU与具体Path不固定绑定"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")
    text(slide, .66, .47, 14.7, .60,
         "Once per layer：动态选择 QoS Path", 28, INK, True)
    panel(slide, .70, 0)
    panel(slide, 8.55, 1)
    shape(slide, 7.99, 1.51, .006, 5.72, "DDDDDD")
    text(slide, .70, 7.61, 14.60, .40,
         "先按请求类别确定 Path 池，再在池内按层动态选路；NPU 与具体 Path 不固定绑定。",
         17, INK)
    text(slide, .70, 8.06, 14.6, .26,
         "方块：Path 中已有的 I/O；箭头：本层待下发的 I/O。方块数量仅作示意，实际压力还与 QoS 服务份额有关。",
         11, GRAY)
    text(slide, .70, 8.42, 14.6, .26,
         "箭头仅示意部分 I/O；同层可使用多条允许的 Path，各 NPU 独立推进。数据所在 ASU 不变。",
         11, GRAY)
    slide.notes_slide.notes_text_frame.text = (
        "概念示意，不是真实仿真数据；四个IO方块表示较高队列占用，一个表示较低占用。\n"
        "方块尺寸与颜色一致，表示同大小的已入Path未完成IO；箭头为本层即将下发的IO，"
        "尚未计入方块数量。本示例假定可比服务条件，实际压力还取决于CIR/PIR及活动状态，"
        "不是断言仅凭4与1即可精确预测等待时间。图中不单独展开正在服务和等待服务的IO。\n"
        "Once per layer对每个request-layer-ASU组合规划一次本层各IO的路径，"
        "不把NPU固定绑定到某个Path，也不表示一层只用一条Path。\n"
        "图中两个NPU当前处理的请求均为SS，故只能在各自目标ASU的SS允许Path集合中选路。"
        "类别由请求参数确定，不是NPU的永久属性；同一NPU后续处理其他类别请求时使用对应池。\n"
        "SS类可选Path框是逻辑集合，不是单独的硬件组。当前配置SS共有96条Path，"
        "跨8个硬件组；这里只展示实际合法的Path0/1/2。SL/LS/LL集合折叠显示，"
        "即使其Path空闲，当前SS请求也不能在该配置下跨类别选用。\n"
        "两图表示各NPU在不同层上的可能选择，不要求所有NPU同步换层。"
        "每层重新规划不等于必须更换路径，图中NPU1至ASU1的Path2保持不变。\n"
        "左图NPU0和NPU1都选择ASU0.Path0，展示Path共享。右图分别选择Path2、Path1。"
        "展示的Path0/1/2属于两个SS请求共同允许的集合；Path共享不等于跨类别任意选路。\n"
        "ASU0.Path0与ASU1.Path0是不同设备上的独立队列。选路不迁移数据、"
        "不改变目标ASU，也不会把已经入队的IO随下一幅图迁走。\n"
        "压力来自客户端可查询的采集状态；没有增加按层实时设备读取、CIR写入、"
        "全局预留账本或NPU重绑定，不混入New Once及后续新策略的能力。\n"
        "蓝色实线来自NPU0，灰色虚线来自NPU1；所有线条、框和文字可编辑。\n"
        "代码依据：shared_path_once.py；policy_logic.py中的允许Path集合及逐层IO规划。"
    )
    target = OUT / "once_per_layer_concept.pptx"
    prs.save(target)
    check = Presentation(target)
    assert len(check.slides) == 1
    assert len(ROUTES[0]) == len(ROUTES[1]) == 4
    assert sum(asu == 0 and path == 0 for _, asu, path in ROUTES[0]) == 2
    assert (1, 1, 2) in ROUTES[0] and (1, 1, 2) in ROUTES[1]
    all_text = "\n".join(s.text for s in check.slides[0].shapes if s.has_text_frame)
    assert all_text.count("当前请求：SS") == 4
    assert all_text.count("SS 类可选 Path（部分）") == 4
    assert all_text.count("其他类别：SL / LS / LL") == 4
    for stage in (0, 1):
        for asu in (0, 1):
            for path in range(3):
                prefix = f"queued_io_s{stage}_asu{asu}_path{path}_"
                count = sum(s.name.startswith(prefix) for s in check.slides[0].shapes)
                assert count == QUEUE_COUNTS[stage][asu][path]
        assert all(QUEUE_COUNTS[stage][asu][path] == 1 for _, asu, path in ROUTES[stage])
    for s in check.slides[0].shapes:
        assert s.shape_type != 13  # No embedded bitmap artwork.
        assert s.left >= 0 and s.top >= 0
        assert s.left + s.width <= check.slide_width + 2
        assert s.top + s.height <= check.slide_height + 2
    print(f"Created {target}; 1 editable slide, {len(check.slides[0].shapes)} objects")


if __name__ == "__main__":
    build()
