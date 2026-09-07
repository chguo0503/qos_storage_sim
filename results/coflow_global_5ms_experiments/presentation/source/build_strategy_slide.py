"""One editable slide, built from the independently audited 36-case results.

Run with the project's Python (python-pptx is already installed). Presentation
names are display labels only; no policy code or experimental output is changed.
"""

import hashlib
import json
from pathlib import Path
from statistics import mean

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "presentation"
AUDIT = ROOT / "data/formal_results/formal_audit_summary.json"
NAVY = "222222"
INK = "222222"
MUTED = "666666"
TEAL = "007D78"
BLUE = "2D65B0"
AMBER = "A85A22"
BG = "FFFFFF"
FONT = "Noto Sans CJK SC"

POLICIES = [
    ("baseline", "固定单路", "Baseline", "每盘读取固定进入 Path0", "", "8795A3"),
    ("once", "压力感知选路", "Once per layer", "每层查看 QoS Path 的 I/O 压力", "动态选择路径", BLUE),
    ("new_once", "预留感知选路", "New Once", "在 Once 基础上，计入全局", "已安排但未完成的 I/O", "218CA0"),
    ("strategy1", "全局客户端调度", "策略 1", "在新 Once 基础上，选择 NPU", "并控制 I/O 下发时机", TEAL),
    ("strategy2", "盘内协同调度", "策略 2", "在策略 1 基础上，增加盘内 I/O 调度", "保留原生 QoS 约束", "5D70AC"),
    ("strategy3", "跨盘协同调度", "策略 3", "在策略 2 基础上，结合各盘读取进度", "协调同一层 I/O 的优先级", "65519A"),
]


def rgb(value):
    return RGBColor.from_string(value)


def shape(slide, x, y, w, h, fill, radius=False, line=None):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if radius:
        s.adjustments[0] = 0.08
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.7)
    else:
        s.line.fill.background()
    s._element.spPr.append(OxmlElement("a:effectLst"))
    style = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        s._element.remove(style)
    return s


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf._txBody.bodyPr.set("anchorCtr", "0")
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    props = run._r.get_or_add_rPr()
    props.set("lang", "zh-CN")
    ea = OxmlElement("a:ea")
    ea.set("typeface", FONT)
    props.append(ea)
    return box


def metrics():
    audit = json.loads(AUDIT.read_text())
    assert audit["complete_case_count"] == 36
    assert audit["independent_audit_passed"]
    values = {}
    for key, *_ in POLICIES:
        for ssu in (5, 6, 7):
            rows = [r for r in audit["rows"] if r["strategy"] == key and r["num_ssu"] == ssu]
            assert len(rows) == 2
            assert {r["seed"] for r in rows} == {20260906, 20260907}
            assert all(r["all_npus_active_whole_window"] for r in rows)
            denominator = sum(r["slo_denominator"] for r in rows)
            assert denominator == 1408
            values[key, ssu] = (
                100 * mean(r["fixed_window_utilization"] for r in rows),
                100 * sum(r["admission_slo_passed"] for r in rows) / denominator,
            )
    return audit, values


def build():
    audit, values = metrics()
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    prs.core_properties.title = "QoS 调度策略对比"
    prs.core_properties.subject = "32 NPU，5/6/7 SSU，同输入，36组正式实验"
    prs.core_properties.author = "QoS Storage Simulation"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    # Conventional technical slide: a plain title and one comparison table.
    text(slide, 0.65, 0.47, 14.7, 0.60, "QoS 调度策略对比", 28, INK, True)
    text(slide, 0.67, 1.16, 14.6, 0.28,
         "32 NPU，单盘 40 GiB/s，状态采集周期 5 ms", 13.5, MUTED)

    x0, y0, tw, header_h, row_h = 0.65, 1.77, 14.70, 0.78, 0.73
    strategy_w, capability_w = 2.73, 4.62
    numeric_x = x0 + strategy_w + capability_w
    group_w = (tw - strategy_w - capability_w) / 3
    shape(slide, x0, y0, tw, header_h, "F0F1F3")
    text(slide, x0 + 0.16, y0 + 0.20, 2.50, 0.35, "策略", 16, INK, True)
    text(slide, x0 + strategy_w + 0.12, y0 + 0.20, 4.3, 0.35, "主要能力", 16, INK, True)
    for j, ssu in enumerate((5, 6, 7)):
        xx = numeric_x + j * group_w
        text(slide, xx, y0 + 0.07, group_w, 0.30, f"{ssu} SSU", 16, INK, True, PP_ALIGN.CENTER)
        text(slide, xx, y0 + 0.46, group_w / 2, 0.23, "利用率 %", 11.5, MUTED, align=PP_ALIGN.CENTER)
        text(slide, xx + group_w / 2, y0 + 0.46, group_w / 2, 0.23, "SLO达标 %", 11.5, MUTED, align=PP_ALIGN.CENTER)

    # Bold per-column maxima, not a preselected favorite strategy.
    maxima = {(ssu, k): max(values[key, ssu][k] for key, *_ in POLICIES)
              for ssu in (5, 6, 7) for k in (0, 1)}
    for i, (key, label, old, line1, line2, accent) in enumerate(POLICIES):
        yy = y0 + header_h + i * row_h
        text(slide, x0 + 0.16, yy + 0.11, strategy_w - 0.26, 0.32, label, 17, INK)
        text(slide, x0 + 0.17, yy + 0.45, strategy_w - 0.27, 0.20, old, 10.5, MUTED)
        cx = x0 + strategy_w + 0.12
        if line2:
            text(slide, cx, yy + 0.105, capability_w - 0.22, 0.26, line1, 14, INK)
            text(slide, cx, yy + 0.395, capability_w - 0.22, 0.26, line2, 14, INK)
        else:
            text(slide, cx, yy + 0.22, capability_w - 0.22, 0.29, line1, 14, INK)
        for j, ssu in enumerate((5, 6, 7)):
            xx = numeric_x + j * group_w
            for k, number in enumerate(values[key, ssu]):
                best = abs(number - maxima[ssu, k]) < 1e-9
                text(slide, xx + k * group_w / 2, yy + 0.19,
                     group_w / 2, 0.35, f"{number:.2f}", 18,
                     INK, best, PP_ALIGN.CENTER, font="DejaVu Sans")
        shape(slide, x0, yy + row_h - 0.006, tw, 0.006, "D6D6D6")
    for xx in (x0 + strategy_w, numeric_x, numeric_x + group_w, numeric_x + 2 * group_w):
        shape(slide, xx, y0, 0.005, header_h + 6 * row_h, "DDDDDD")

    text(slide, x0 + 0.02, 7.30, tw - 0.04, 0.32,
         "策略 1–3 的接纳后 SLO 达标率提高，但整批请求完成时间均长于 Once。", 15, INK)
    footnotes = [
        "注：利用率取 [1,2] 秒内 32 卡平均；SLO 按接纳至 8 层完成 ≤ 1.5 × 纯计算总时间统计，不含接纳前排队。",
        "两次实验汇总，每次 704 请求；静态 CIR，控制开销未计入。粗体为列内最高值。盘内及跨盘调度需要盘侧配合。",
    ]
    for i, line in enumerate(footnotes):
        text(slide, x0 + 0.02, 8.00 + i * 0.28, tw - 0.04, 0.24, line, 11, MUTED)

    slide.notes_slide.notes_text_frame.text = (
        "展示名称仅用于汇报，没有修改任何策略源码或重新选择实验结果。\n"
        "名称映射：" + "；".join(f"{old}={label}" for _, label, old, *_ in POLICIES) + "。\n"
        "New Once沿用Once的选路算法，但压力计数使用所有受管流量的全局已规划减HBM完成账本，"
        "不是在旧SSD快照上简单加一遍，也不是增加读取SSD状态频率。\n"
        "S1新增新请求到达时的NPU绑定与全局未下发I/O的授予/延迟；不移动已经存储的数据。\n"
        "S2可改变同Path尚未服务队列头，并在原生最小虚拟完成时间+EPS等价候选内选择跨Path服务；"
        "保持CIR/WFQ计费，不是任意EDF或全局最优，不抢占正在服务的I/O。\n"
        "S3在S2权限上使用同层各SSD已激活未HBM完成数据的共同进度选择优先级；不保证所有盘同时服务同一层。\n"
        "6 SSU客户端参数只用seed20260906选择；seed20260907以及其他盘数为泛化检查。"
        "两种子平均窗口利用率；两种子合并SLO分母1408。SLO门限为1.5倍该请求八层纯计算总时长。\n"
        "5盘持续输入超过容量，但有限请求可排空；不是仅由平均输入证明[1,2]秒利用率上限。"
        "S1–S3所有拓扑/种子的整批完成时间都比Once更长，0/6组更快；不能声称整批吞吐提升。\n"
        "所有图形、文字、结果数字都是可编辑PPT对象；无栅格化表格。\n"
        f"来源：{AUDIT.relative_to(ROOT)}\n"
        f"来源SHA256：{hashlib.sha256(AUDIT.read_bytes()).hexdigest()}\n"
        "完整报告：../coflow_global_5ms_report.pdf\n"
        + json.dumps({f"{key}_ssu{ssu}": list(value) for (key, ssu), value in values.items()}, ensure_ascii=False)
    )
    OUT.mkdir(exist_ok=True)
    target = OUT / "coflow_strategy_summary.pptx"
    prs.save(target)

    # Check editability, one-slide constraint, numeric content and geometry.
    check = Presentation(target)
    assert len(check.slides) == 1
    assert not any(s.shape_type == 13 for s in check.slides[0].shapes)
    all_text = "\n".join(s.text for s in check.slides[0].shapes if s.has_text_frame)
    for pair in values.values():
        for number in pair:
            assert f"{number:.2f}" in all_text
    for s in check.slides[0].shapes:
        assert s.left >= 0 and s.top >= 0
        assert s.left + s.width <= check.slide_width + 2
        assert s.top + s.height <= check.slide_height + 2
    manifest = {
        "slide_count": 1,
        "all_text_and_shapes_editable": True,
        "source": str(AUDIT.relative_to(ROOT)),
        "source_sha256": hashlib.sha256(AUDIT.read_bytes()).hexdigest(),
        "pptx_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "formal_cases": audit["complete_case_count"],
        "name_mapping": {key: label for key, label, *_ in POLICIES},
        "metrics": {f"{key}_ssu{ssu}": {"utilization_pct": pair[0], "admission_slo_pct": pair[1]}
                    for (key, ssu), pair in values.items()},
    }
    (OUT / "source/slide_audit.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"pptx": str(target), "slides": 1, "shapes": len(check.slides[0].shapes)}, ensure_ascii=False))


if __name__ == "__main__":
    build()
