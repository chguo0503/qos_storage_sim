"""One editable page: routing concept plus verified warm experiment results."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_strategy_slide import OUT, EXPERIMENT, box, label, route, segment, storage


def build():
    original = OUT / 'fixed_candidate_pool_once.pptx'
    original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
    prs = Presentation(original)
    slide = prs.slides[0]
    # Keep the existing theme and explanatory notes, rebuild only the new copy.
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)

    rows = list(csv.DictReader((EXPERIMENT/'macro_summary.csv').open()))
    stats = {(r['scenario'], r['policy']): r for r in rows
             if r['window'] == 'warm_2_4s' and r['policy'] in ('baseline', 'static')}
    assert len(stats) == 4
    assert all(int(r['seed_count']) == 3 for r in stats.values())

    label(slide, '固定候选池 + Once：先限定范围，再按层选路',
          .48, .23, 15.05, .62, size=28, bold=True)
    label(slide,
          '本次24种画像的已接纳请求：每层计算 C > 12 ms → 固定8条/盘；否则 → 原类别全池。',
          .58, 1.05, 15., .35, size=16.)
    label(slide,
          '初始每层等待余量 = C / 2。橙底8条由同盘同类的收窄请求共享；其余Path仍供全池请求使用。',
          .58, 1.47, 15., .26, size=12.5, color='gray')
    segment(slide, 7.99, 1.85, 7.99, 5.22, color='divider', width=.65)

    for dx, heading, queues in ((0., '第 L 层', (4, 1, 2)),
                                (7.85, '第 L+1 层', (1, 3, 4))):
        label(slide, heading, .70+dx, 1.86, 2.3, .42, size=20, bold=True)
        label(slide, '同一块SSU；候选池不变，选路可变',
              3.35+dx, 1.92, 4.05, .27, size=12.0, color='gray',
              align=PP_ALIGN.RIGHT)
        tx, mids = storage(slide, dx, 2.43, 0, queues)
        for number, yy, narrow in ((0, 2.62, False), (1, 4.02, True)):
            color = 'orange' if narrow else 'blue'
            box(slide, .76+dx, yy, 1.61, .58, line=color, width=.85)
            label(slide, f'NPU{number}', .76+dx, yy, 1.61, .58,
                  size=18, color=color, bold=True, align=PP_ALIGN.CENTER)
            label(slide,
                  'SL · 32K/miss4K · C=28.59ms' if narrow
                  else 'SL · 32K/miss1K · C=7.26ms',
                  .34+dx, yy+.69, 2.46, .27, size=10.9,
                  color='gray', align=PP_ALIGN.CENTER)
            label(slide, '固定池：8条/盘' if narrow else '保留全池：32条/盘',
                  .44+dx, yy+.99, 2.25, .27, size=12., color=color,
                  bold=True, align=PP_ALIGN.CENTER)
        if not dx:
            route(slide, 2.37, 2.91, 3.03, tx, mids[13], color='blue')
            route(slide, 2.37, 4.31, 3.66, tx, mids[44], color='orange', dashed=True)
        else:
            route(slide, 2.37+dx, 2.91, 3.03+dx, tx, mids[12]-.065, color='blue')
            route(slide, 2.37+dx, 4.31, 3.66+dx, tx, mids[12]+.065,
                  color='orange', dashed=True)

    label(slide,
          '蓝线：全池请求；橙虚线：固定池请求。Once用每5ms更新的拥塞快照逐块选路；一层可走多条Path。',
          .61, 5.39, 14.95, .31, size=13.5)
    label(slide,
          '实测结果：32 NPU / 3 SSU × 40 GiB/s · Random输入 · warm [2,4)秒 · seed 7 / 19 / 43 等权平均',
          .61, 5.87, 14.95, .31, size=13.5, bold=True)

    # Native shapes/text rather than a screenshot: every result stays editable.
    x0, y0 = .61, 6.30
    widths = (2.85, 4.20, 3.30, 4.50)
    row_h = .365
    header = ('负载', '策略', 'NPU平均利用率', 'TTFT SLO×1.5达标率')
    display = [header]
    for scenario, title in (('semi', '间歇过载'), ('full', '持续过载')):
        for policy, policy_title in (('baseline', 'Baseline Random'),
                                     ('static', '固定候选池 + Once')):
            r = stats[(scenario, policy)]
            display.append((title, policy_title, f"{float(r['mean_U_percent']):.2f}%",
                            f"{float(r['mean_slo_percent']):.2f}%"))
    for i, values in enumerate(display):
        x = x0
        for j, (w, value) in enumerate(zip(widths, values)):
            fill = 'io' if i == 0 else ('orange_bg' if i in (2, 4) else 'white')
            box(slide, x, y0+i*row_h, w, row_h, fill=fill, line='line', width=.5)
            label(slide, value, x+.10, y0+i*row_h+.012, w-.20, row_h-.024,
                  size=14., bold=(i == 0 or (i in (2, 4) and j >= 2)),
                  align=PP_ALIGN.LEFT if j < 2 else PP_ALIGN.CENTER)
            x += w

    label(slide,
          'SLO口径：窗内接纳请求跟踪至完成；接纳到prefill完成 ≤ 1.5 × 自身纯计算时间，不含接纳前排队。',
          .61, 8.27, 14.95, .26, size=11.5, color='gray')
    ll_base = float(stats[('full', 'baseline')]['mean_LL_slo_percent'])
    ll_static = float(stats[('full', 'static')]['mean_LL_slo_percent'])
    label(slide,
          f'代价：持续过载的LL类SLO达标率 {ll_base:.2f}% → {ll_static:.2f}%。上图仅示意2卡、1盘；未接纳首层用全池；FIFO、CIR/PIR不变。',
          .61, 8.62, 14.95, .26, size=11.5, color='gray')

    notes = slide.notes_slide.notes_text_frame
    notes.text = notes.text.replace(
        '图中仅示意两张NPU与两块SSU，第三块同理。',
        '图中仅示意两张NPU与一块SSU，其他盘同理。')
    notes.text += '\n\n本版补充正式实测对照表，仍只有1页。\n'
    notes.text += '数据来源：../macro_summary.csv，window=warm_2_4s，policy=baseline/static；与逐种子comparison.csv独立复核一致。\n'
    notes.text += '新版正文为避免拥挤只画一块SSU，实验本身仍是32NPU、3SSU。图中队列长度和箭头仅解释机制，并非日志采样。\n'
    notes.text += 'NPU利用率：warm [2,4)内实际计算卡时间/(32×2秒)。SLO样本：在该窗口接纳的全部请求，跟踪至最后完成，包括窗口结束后才完成的请求。\n'
    notes.text += '每个种子先计算比例，再三个种子等权平均；两策略的请求集合与卡内输入顺序按种子完全一致。\n'
    notes.text += '本次的TTFT标签沿用实验叫法，实质为接纳至prefill完成，不是从外部到达至真实首token。\n'
    for scenario in ('semi', 'full'):
        for policy in ('baseline', 'static'):
            r = stats[(scenario, policy)]
            notes.text += f"{scenario}/{policy}: U={r['mean_U_percent']}%, SLO={r['mean_slo_percent']}%\n"

    target = OUT/'fixed_candidate_pool_once_with_results.pptx'
    prs.core_properties.title = '固定候选池 + Once：策略与实测结果'
    prs.save(target)
    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_hash
    reopened = Presentation(target)
    assert len(reopened.slides) == 1
    outside = [s.name for s in reopened.slides[0].shapes
               if s.left < 0 or s.top < 0 or s.left+s.width > Inches(16)
               or s.top+s.height > Inches(9)]
    assert not outside, outside
    check = {'slides': 1, 'source_ppt_unchanged': True, 'outside_shapes': outside,
             'data_source': '../macro_summary.csv', 'window': 'warm_2_4s',
             'seeds': [7, 19, 43], 'table': display,
             'pptx_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
             'visual_review': 'pending'}
    (OUT/'slide_with_results_checks.json').write_text(
        json.dumps(check, ensure_ascii=False, indent=2)+'\n')
    print(target)


if __name__ == '__main__':
    build()
