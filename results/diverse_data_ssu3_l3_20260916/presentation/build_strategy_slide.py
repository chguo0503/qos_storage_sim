"""Build one editable strategy slide, following the reference's third page.

No simulations are run. The queue lengths and arrows are explanatory drawings,
not samples from the experiment. Source data and the source PPT are read-only.
"""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
EXPERIMENT = OUT.parent
REFERENCE = Path('/home/chguo/work/once_per_layer_concept.pptx')
FONT = 'Noto Sans CJK SC'
C = {
    'ink': '222222', 'gray': '666666', 'line': 'B4C0D0',
    'blue': '3569A4', 'orange': 'B46620', 'orange_bg': 'FFF1DF',
    'white': 'FFFFFF', 'io': 'E4E9F0', 'soft': 'F7F9FB',
    'divider': 'DDDDDD',
}


def rgb(name):
    return RGBColor.from_string(C.get(name, name))


def box(slide, x, y, w, h, *, fill='white', line='line', width=.65, name=''):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or f'box_{len(slide.shapes)}'
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape


def label(slide, text, x, y, w, h, *, size=16, color='ink', bold=False,
          align=PP_ALIGN.LEFT, name=''):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or text[:48]
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    # Keep both Chinese and Latin editable and consistent in PowerPoint/LO.
    for tag in ('a:ea', 'a:cs'):
        node = OxmlElement(tag)
        node.set('typeface', FONT)
        r._r.get_or_add_rPr().append(node)
    return shape


def segment(slide, x0, y0, x1, y1, *, color, dashed=False, arrow=False,
            width=1.55):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                  Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    s._element.spPr.append(OxmlElement('a:effectLst'))
    ln = s._element.spPr.get_or_add_ln()
    if dashed:
        dash = OxmlElement('a:prstDash')
        dash.set('val', 'dash')
        ln.append(dash)
    if arrow:
        head = OxmlElement('a:tailEnd')
        head.set('type', 'triangle')
        head.set('w', 'sm')
        head.set('len', 'sm')
        ln.append(head)
    return s


def route(slide, sx, sy, bend_x, tx, ty, *, color, dashed=False):
    segment(slide, sx, sy, bend_x, sy, color=color, dashed=dashed)
    segment(slide, bend_x, sy, bend_x, ty, color=color, dashed=dashed)
    segment(slide, bend_x, ty, tx, ty, color=color, dashed=dashed, arrow=True)


def storage(slide, dx, y, disk, lengths):
    """A subset of the actual SL legal path IDs: 12/44 are in the fixed pool."""
    x = 4.20 + dx
    box(slide, x, y, 3.20, 2.48, name=f'panel_{dx}_SSU{disk}')
    label(slide, f'SSU{disk}', x+.16, y+.08, 1.10, .30, size=16.5, bold=True)
    label(slide, 'SL 类：32 条 Path（仅画 3 条）',
          x+.16, y+.45, 2.90, .27, size=11.5, bold=True)
    mids = {}
    for row, (pid, count) in enumerate(zip((12, 13, 44), lengths)):
        yy = y+.82+row*.49
        narrow = pid in (12, 44)
        box(slide, x+.13, yy, 2.94, .43,
            fill='orange_bg' if narrow else 'white',
            line='orange' if narrow else 'line', width=.60,
            name=f'panel_{dx}_ssu_{disk}_path_{pid}_candidate_{narrow}')
        label(slide, f'Path{pid}', x+.27, yy+.015, .97, .39,
              size=14.0, color='ink')
        for j in range(count):
            bx = x+1.26+j*.39
            box(slide, bx, yy+.065, .32, .30,
                fill='io', line='line', width=.55)
            label(slide, 'IO', bx, yy+.065, .32, .30, size=9.5,
                  align=PP_ALIGN.CENTER)
        mids[pid] = yy+.215
    label(slide, '每盘独立选路 · Path 内部 FIFO',
          x+.16, y+2.285, 2.90, .17, size=9.8, color='gray',
          align=PP_ALIGN.CENTER)
    return x+.13, mids


def npu(slide, dx, y, n, compute_ms, miss_k, narrow):
    color = 'orange' if narrow else 'blue'
    box(slide, .76+dx, y, 1.61, .64, line=color, width=.85)
    label(slide, f'NPU{n}', .76+dx, y, 1.61, .64,
          size=18, color=color, bold=True, align=PP_ALIGN.CENTER)
    # Both examples are SL: the new restriction is not the old SS/SL partition.
    label(slide, f'SL · 总长32K / miss{miss_k}K',
          .44+dx, y+.75, 2.25, .26, size=11.2,
          color='gray', align=PP_ALIGN.CENTER)
    label(slide, f'每层计算 {compute_ms:.2f} ms',
          .44+dx, y+1.03, 2.25, .26, size=12.0,
          align=PP_ALIGN.CENTER)
    label(slide, '固定池：8条/盘' if narrow else '保留全池：32条/盘',
          .44+dx, y+1.31, 2.25, .26, size=12.0, bold=True,
          color=color, align=PP_ALIGN.CENTER)


def build():
    source_hash = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    rows = list(csv.DictReader((EXPERIMENT/'input_profile_summary.csv').open()))
    profiles = {int(r['miss']): r for r in rows if int(r['total_K']) == 32}
    short_c = float(profiles[1024]['layer_C_ms'])
    long_c = float(profiles[4096]['layer_C_ms'])
    assert profiles[1024]['category'] == profiles[4096]['category'] == 'SL'
    assert profiles[1024]['static_pool'] == 'full'
    assert profiles[4096]['static_pool'] == 'narrow8'

    # Reuse the actual reference's theme/master, remove its pages in memory only.
    prs = Presentation(REFERENCE)
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for s in list(slide.shapes):
        s._element.getparent().remove(s._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb('white')
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)

    label(slide, '固定候选池 + Once：先限定范围，再按层选路',
          .48, .23, 15.05, .62, size=28, bold=True)
    label(slide,
          '本次24种画像的已接纳请求：每层计算 C > 12 ms → 固定8条/盘；否则 → 原类别全池（本例 SL 为32条/盘）。',
          .58, 1.07, 15.0, .36, size=16.0)
    label(slide,
          'SLO = 1.5 × 纯计算时间 → 初始每层等待余量 = C / 2；C 越长，初始可等待的时间也越长。',
          .58, 1.48, 15.0, .24, size=12.4, color='gray')
    segment(slide, 7.99, 1.78, 7.99, 7.39, color='divider', width=.65)

    for dx, text, queues in (
        (0., '第 L 层', ((4, 1, 2), (2, 4, 1))),
        (7.85, '第 L+1 层', ((1, 3, 4), (4, 1, 2))),
    ):
        label(slide, text, .70+dx, 1.79, 2.30, .42, size=20, bold=True)
        label(slide, '候选池不变，重新按拥塞选路' if dx else '先确定候选池，再按拥塞选路',
              3.13+dx, 1.86, 4.27, .27, size=12.5, color='gray',
              align=PP_ALIGN.RIGHT)
        tx0, m0 = storage(slide, dx, 2.18, 0, queues[0])
        tx1, m1 = storage(slide, dx, 4.90, 1, queues[1])
        npu(slide, dx, 3.07, 0, short_c, 1, False)
        npu(slide, dx, 5.77, 1, long_c, 4, True)
        sx = 2.37+dx
        if not dx:
            # Solid blue: full pool, so both white and orange paths are legal.
            route(slide, sx, 3.24, 3.02, tx0, m0[13], color='blue')
            route(slide, sx, 3.54, 2.85, tx1, m1[44], color='blue')
            # Dashed orange: shared fixed subset only.
            route(slide, sx, 5.94, 3.69, tx0, m0[44], color='orange', dashed=True)
            route(slide, sx, 6.23, 3.48, tx1, m1[12], color='orange', dashed=True)
        else:
            route(slide, sx, 3.24, 3.02+dx, tx0, m0[12]-.07, color='blue')
            route(slide, sx, 3.54, 2.85+dx, tx1, m1[13], color='blue')
            route(slide, sx, 5.94, 3.69+dx, tx0, m0[12]+.07,
                  color='orange', dashed=True)
            route(slide, sx, 6.23, 3.48+dx, tx1, m1[44],
                  color='orange', dashed=True)

    box(slide, .61, 7.69, .23, .23, fill='orange_bg', line='orange')
    label(slide,
          '橙底8条：每组固定1条，同盘同类的收窄请求共享；白底其余24条：仍供全池请求使用。',
          .94, 7.63, 14.6, .35, size=15.0)
    label(slide,
          'Once 用每5ms更新的最新拥塞快照，在池内逐块规划本层I/O；一层可走多条Path，固定的是候选范围。',
          .61, 8.09, 14.95, .36, size=16.0)
    label(slide,
          '32 NPU / 3 SSU，仅画2卡2盘示意。目标是改善整体SLO，可能牺牲部分请求；未接纳首层仍用全池；CIR/PIR不变。',
          .61, 8.58, 14.95, .25, size=11.8, color='gray')

    slide.notes_slide.notes_text_frame.text = '''固定候选池 + Once（static）

参考：/home/chguo/work/once_per_layer_concept.pptx，第3页。
沿用左右第L层/第L+1层、NPU到SSU折线、Path队列方块的表达。
本页为说明机制的示意图，IO队列长度与箭头不是实测日志；只画部分Path和部分I/O。
实验实际配置：32 NPU、3 SSU，每盘40 GiB/s；图中仅示意两张NPU与两块SSU，第三块同理。

1. 新增的动作：上层使用本请求初始画像，决定候选Path范围。
C为每层纯计算时间。8层请求的纯计算时间为8C，实验SLO预算为1.5×8C，初始等待预算为4C。
均摊每层初始等待余量=(1.5×8C−8C)/8=C/2。
通用收窄条件：C/2 > max(6ms, 本层独占读取下界)。6ms是当前实验启发式参数。
本次24画像的独占读取下界均小于6ms，所以本批数据可以简化为C>12ms；不是C>6ms，也不是通用阈值。
独占读取下界(ms)=1000×max(max_d(V_d/40),sum_d(V_d)/50)，V_d单位GiB。
它只是假设无他人竞争时的乐观物理下界，不代表实际独占，也不是实际完成时间预测。
static每层沿用初始预算，不随截止临近而恢复全池；不能理解为动态判断实际紧急程度。

2. 原类别池与固定池。
每盘256条Path，8个硬件组，每组32条。组内SS/SL/LS/LL各12/4/12/4条。
原类别全池分别96/32/96/32条。固定池从每组取该类别的第一条，合计8条。
本页特意使用两个同属SL的真实data画像，避免把新增策略误认为原SS/SL分类。
SL每盘32条候选Path；固定8条的全局ID为12,44,76,108,140,172,204,236。
图中橙色Path12和Path44均为固定池成员，白色Path13是其余24条中的一条。
NPU0：总长32K，miss1024，每层计算7.257231579079756ms，保留全池32条。
NPU1：总长32K，miss4096，每层计算28.592841995880782ms，收窄为8条。
数据来自../input_profile_summary.csv；总长包含miss部分，不额外相加。
所有同盘同类收窄请求共用同一组8条，不是每NPU各自独享8条，也不是全系统总共只有8条。
其余24条没有关闭、没有改CIR/PIR，已有IO正常服务，其他同类全池请求仍可使用。

3. 原Once部分保持。
盘侧每5ms采样的Path拥塞计数与已知QoS配置，输入原Once预计完成时间计算。
每个请求/层/SSU调用一次规划；调用内部逐IO块选路，并更新本地影子计数。一层可以跨多条Path。
图中蓝色实线为全池请求，能选择橙底或白底Path；橙色虚线为收窄请求，只能选择橙底Path。
第L+1层候选范围不变，但队列变化，选择的具体Path可以变化。
上层决定候选池，盘侧提供周期拥塞快照；Path内的已提交IO继续FIFO。
未接纳的跨请求首层预取没有接纳时钟预算，使用原Once全池；本页主体画已接纳请求。

4. 为什么可能改善SLO，以及边界。
把初始等待余量较大的请求集中到较少的加权Path，可能让其他请求获得更多服务机会。
本策略没有改变请求顺序，没有改FIFO，没有改变CIR/PIR，没有固定预留带宽，也不保证实际供给b_i达到需求B_i=V_i/C_i。
候选池限制只是启发式服务机会调整，不保证每个请求TTFT达标，也不保证所有类别或NPU利用率都上升。
实验主要SLO为“接纳到prefill完成 <=1.5×纯计算时间”，不包含接纳前排队，也没有实际首token事件。
已完成的持续过载3种子warm均值：baseline SLO40.30%，static76.96%；LL类别反而95.08%降至27.17%。
因此“改善整体SLO”是目标与本次总体结果，不能表述成兼顾所有请求的无代价保证。

实现核对：../policy.py中的upper_budget、choose_path_pool、install_policy。
实验口径：../METHOD.md。数据与结果：../input_profile_summary.csv、../README.md。
'''

    prs.core_properties.title = '固定候选池 + Once：先限定范围，再按层选路'
    prs.core_properties.subject = 'L3 fixed candidate pool routing, one-slide concept'
    prs.core_properties.author = 'QoS Storage Simulation'
    prs.core_properties.keywords = 'Once per layer, QoS Path, SLO, static'
    pptx = OUT/'fixed_candidate_pool_once.pptx'
    prs.save(pptx)

    # Artifact checks: source stays untouched, one editable page, no raster shapes.
    reopened = Presentation(pptx)
    assert len(reopened.slides) == 1
    bounds = []
    for s in reopened.slides[0].shapes:
        if s.left < 0 or s.top < 0 or s.left+s.width > prs.slide_width or s.top+s.height > prs.slide_height:
            bounds.append(s.name)
    assert not bounds, bounds
    assert all(s.shape_type != 13 for s in reopened.slides[0].shapes)
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == source_hash
    check = dict(
        reference=str(REFERENCE), reference_page=3, reference_sha256=source_hash,
        reference_unchanged=True, output=pptx.name, slide_count=1,
        dimensions_in=[16, 9], all_shapes_editable=True,
        shapes=len(reopened.slides[0].shapes), out_of_slide_bounds=bounds,
        illustrative_not_measured=True,
        example_profiles=[profiles[1024], profiles[4096]],
        preview_review='pending',
    )
    (OUT/'slide_checks.json').write_text(json.dumps(check, ensure_ascii=False, indent=2)+'\n')
    print(pptx)


if __name__ == '__main__':
    build()
