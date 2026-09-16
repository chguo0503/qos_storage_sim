"""One editable page: mechanism plus Baseline / original Once / fixed-pool data.

Requires complete three-seed warm summaries for both scenarios. Missing Once
controls are an error; the script never exports invented values/placeholders.
Only a new PPT and its own check JSON are written; original artifacts stay intact.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_strategy_slide import OUT, EXPERIMENT, box, label, route, segment, storage


def build():
    original = OUT / 'fixed_candidate_pool_once.pptx'
    original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
    prs = Presentation(original)
    slide = prs.slides[0]
    # Keep the existing theme and explanatory notes, rebuild only the new copy.
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)

    baseline_source = EXPERIMENT/'macro_summary.csv'
    once_source = EXPERIMENT/'once_control_macro_summary.csv'
    if not once_source.is_file():
        raise FileNotFoundError(
            f"Original Once controls are not ready: {once_source}. "
            "Complete semi/full × seeds 7/19/43 before building this PPT.")
    source_hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                     for p in (baseline_source, once_source)}
    rows = list(csv.DictReader(baseline_source.open()))
    once_rows = list(csv.DictReader(once_source.open()))
    selected = [r for r in rows if r['window']=='warm_2_4s'
                and r['scenario'] in ('semi','full') and r['policy'] in ('baseline','static')]
    selected += [r for r in once_rows if r['window']=='warm_2_4s'
                 and r['scenario'] in ('semi','full') and r['policy']=='once']
    stats = {(r['scenario'],r['policy']):r for r in selected}
    expected = {(s,p) for s in ('semi','full') for p in ('baseline','once','static')}
    if set(stats) != expected or len(selected) != len(expected):
        raise ValueError(f"Need exactly six warm summary rows; got {sorted(stats)} "
                         f"from {len(selected)} selected rows")
    for key,r in stats.items():
        seeds = json.loads(r['seeds'])
        if int(r['seed_count']) != 3 or len(seeds) != 3 or set(seeds) != {7,19,43}:
            raise ValueError(f"Incomplete seed controls for {key}: {r['seeds']}")
        for field in ('mean_U_percent','mean_slo_percent'):
            value = float(r[field])
            if not math.isfinite(value) or not 0 <= value <= 100:
                raise ValueError(f"Invalid {field} for {key}: {r[field]}")

    label(slide, '固定候选池 + Once：先限定范围，再按层选路',
          .48, .23, 15.05, .62, size=28, bold=True)
    label(slide,
          '本次24种画像的已接纳请求：每层计算 C > 12 ms → 固定8条/盘；否则 → 原类别全池。',
          .58, 1.05, 15., .35, size=16.)
    label(slide,
          '初始每层等待余量 = C / 2。橙底8条由同盘同类的收窄请求共享；其余Path仍供全池请求使用。',
          .58, 1.47, 15., .26, size=12.5, color='gray')
    segment(slide, 7.99, 1.85, 7.99, 5.22, color='divider', width=.65)

    for dx, heading, queues in ((0., '第 L 层', (4, 1, 2)),
                                (7.85, '第 L+1 层', (1, 3, 4))):
        label(slide, heading, .70+dx, 1.86, 2.3, .42, size=20, bold=True)
        label(slide, '同一块SSU；候选池不变，选路可变',
              3.35+dx, 1.92, 4.05, .27, size=12.0, color='gray',
              align=PP_ALIGN.RIGHT)
        tx, mids = storage(slide, dx, 2.43, 0, queues)
        for number, yy, narrow in ((0, 2.62, False), (1, 4.02, True)):
            color = 'orange' if narrow else 'blue'
            box(slide, .76+dx, yy, 1.61, .58, line=color, width=.85)
            label(slide, f'NPU{number}', .76+dx, yy, 1.61, .58,
                  size=18, color=color, bold=True, align=PP_ALIGN.CENTER)
            label(slide,
                  'SL · 32K/miss4K · C=28.59ms' if narrow
                  else 'SL · 32K/miss1K · C=7.26ms',
                  .34+dx, yy+.69, 2.46, .27, size=10.9,
                  color='gray', align=PP_ALIGN.CENTER)
            label(slide, '固定池：8条/盘' if narrow else '保留全池：32条/盘',
                  .44+dx, yy+.99, 2.25, .27, size=12., color=color,
                  bold=True, align=PP_ALIGN.CENTER)
        if not dx:
            route(slide, 2.37, 2.91, 3.03, tx, mids[13], color='blue')
            route(slide, 2.37, 4.31, 3.66, tx, mids[44], color='orange', dashed=True)
        else:
            route(slide, 2.37+dx, 2.91, 3.03+dx, tx, mids[12]-.065, color='blue')
            route(slide, 2.37+dx, 4.31, 3.66+dx, tx, mids[12]+.065,
                  color='orange', dashed=True)

    label(slide,
          '蓝线：全池请求；橙虚线：固定池请求。Once用每5ms更新的拥塞快照逐块选路；一层可走多条Path。',
          .61, 5.39, 14.95, .31, size=13.5)
    label(slide,
          '实测结果：32 NPU / 3 SSU × 40 GiB/s · Random输入 · warm [2,4)秒 · seed 7 / 19 / 43 等权平均',
          .61, 5.81, 14.95, .27, size=13.0, bold=True)

    # Native shapes/text rather than a screenshot: every result stays editable.
    x0, y0 = .61, 6.16
    widths = (2.85, 4.20, 3.30, 4.50)
    row_h = .290
    header = ('负载', '策略', 'NPU平均利用率', 'TTFT SLO×1.5达标率')
    display = [header]
    for scenario, title in (('semi', '间歇过载'), ('full', '持续过载')):
        for policy, policy_title in (('baseline', 'Baseline Random'),
                                     ('once', '原始 Once per layer'),
                                     ('static', '固定候选池 + Once')):
            r = stats[(scenario, policy)]
            display.append((title, policy_title, f"{float(r['mean_U_percent']):.2f}%",
                            f"{float(r['mean_slo_percent']):.2f}%"))
    for i, values in enumerate(display):
        x = x0
        for j, (w, value) in enumerate(zip(widths, values)):
            fill = 'io' if i == 0 else ('orange_bg' if i in (3, 6) else 'white')
            box(slide, x, y0+i*row_h, w, row_h, fill=fill, line='line', width=.5)
            label(slide, value, x+.10, y0+i*row_h+.012, w-.20, row_h-.024,
                  size=13.5, bold=(i == 0 or (i in (3, 6) and j >= 2)),
                  align=PP_ALIGN.LEFT if j < 2 else PP_ALIGN.CENTER)
            x += w

    label(slide,
          'SLO口径：窗内接纳请求跟踪至完成；接纳到prefill完成 ≤ 1.5 × 自身纯计算时间，不含接纳前排队。',
          .61, 8.30, 14.95, .26, size=11.5, color='gray')
    ll_base = float(stats[('full', 'baseline')]['mean_LL_slo_percent'])
    ll_once = float(stats[('full', 'once')]['mean_LL_slo_percent'])
    ll_static = float(stats[('full', 'static')]['mean_LL_slo_percent'])
    label(slide,
          f'持续过载LL达标率（Baseline / 原始Once / 固定池）：{ll_base:.2f}% / {ll_once:.2f}% / {ll_static:.2f}%；示意仅2卡1盘；未接纳首层用全池；FIFO/CIR/PIR不变。',
          .61, 8.65, 14.95, .23, size=11.3, color='gray')

    notes = slide.notes_slide.notes_text_frame
    notes.text = notes.text.replace(
        '图中仅示意两张NPU与两块SSU，第三块同理。',
        '图中仅示意两张NPU与一块SSU，其他盘同理。')
    notes.text += '\n\n本版补充Baseline Random、原始Once per layer、固定候选池+Once三组正式实测，仍只有1页。\n'
    notes.text += '数据来源：../macro_summary.csv（baseline/static）和../once_control_macro_summary.csv（once），window=warm_2_4s，三个种子7/19/43等权平均。\n'
    notes.text += '新版正文为避免拥挤只画一块SSU，实验本身仍是32NPU、3SSU。图中队列长度和箭头仅解释机制，并非日志采样。\n'
    notes.text += 'NPU利用率：warm [2,4)内实际计算卡时间/(32×2秒)。SLO样本：在该窗口接纳的全部请求，跟踪至最后完成，包括窗口结束后才完成的请求。\n'
    notes.text += '每个种子先计算比例，再三个种子等权平均；三策略使用按种子完全相同的完整输入与卡内请求顺序；warm窗口接纳请求集合可能不同。\n'
    notes.text += '本次的TTFT标签沿用实验叫法，实质为接纳至prefill完成，不是从外部到达至真实首token。\n'
    notes.text += '原始Once对每个请求始终在原类别完整合法候选池内，按最近5ms拥塞快照逐块规划本层I/O；没有新增的8条候选池限制。固定候选池+Once仅先缩小部分已接纳请求的候选范围，池内规划引擎仍是同一个Once。Baseline全部I/O使用Path0。三者都不重排Path内部FIFO、不写CIR/PIR。\n'
    notes.text += '机制主图解释固定候选池+Once；原始Once的同类请求均可选择橙底与白底Path。橙底不是专用带宽、不是请求独享。\n'
    for scenario in ('semi', 'full'):
        for policy in ('baseline', 'once', 'static'):
            r = stats[(scenario, policy)]
            notes.text += f"{scenario}/{policy}: U={r['mean_U_percent']}%, SLO={r['mean_slo_percent']}%\n"

    target = OUT/'fixed_candidate_pool_once_threeway.pptx'
    prs.core_properties.title = '固定候选池 + Once：Baseline、原始Once与固定候选池对照'
    prs.save(target)
    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_hash
    reopened = Presentation(target)
    assert len(reopened.slides) == 1
    outside = [s.name for s in reopened.slides[0].shapes
               if s.left < 0 or s.top < 0 or s.left+s.width > Inches(16)
               or s.top+s.height > Inches(9)]
    assert not outside, outside
    check = {'slides': 1, 'source_ppt_unchanged': True, 'outside_shapes': outside,
             'data_sources': source_hashes, 'window': 'warm_2_4s',
             'seeds': [7, 19, 43], 'table': display,
             'pptx_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
             'visual_review': 'pending'}
    (OUT/'slide_threeway_checks.json').write_text(
        json.dumps(check, ensure_ascii=False, indent=2)+'\n')
    print(target)


if __name__ == '__main__':
    build()
